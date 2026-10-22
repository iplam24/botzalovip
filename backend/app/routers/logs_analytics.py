"""Tổng quan, nhật ký, phân tích và xuất báo cáo CSV — mọi số liệu từ DB thật."""
import csv
import io
from datetime import datetime
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from ..database import get_db
from ..models import *
from ..deps import get_current_user

router = APIRouter(tags=["Tổng quan / Nhật ký / Phân tích"])


@router.get("/api/suc-khoe")
def health():
    return {"trang_thai": "khoe", "dich_vu": "Zalo AI Hub"}


@router.get("/api/tong-quan")
def overview(db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    # Trạng thái AI/Zalo thật: provider mặc định + các tài khoản Zalo đã cấu hình
    provider_mac_dinh = db.query(AIProviderConfig).filter(AIProviderConfig.is_default == True).first()
    tai_khoan_zalo = [
        {"id": a.id, "name": a.name, "account_type": a.account_type, "status": a.status}
        for a in db.query(ZaloAccount).all()
    ]
    return {
        "tong_khach_hang": db.query(Customer).count(),
        "tong_hoi_thoai": db.query(Conversation).count(),
        "hoi_thoai_dang_xu_ly": db.query(Conversation).filter(Conversation.status == "dang_xu_ly").count(),
        "hoi_thoai_chuyen_nhan_vien": db.query(Conversation).filter(Conversation.human_takeover == True).count(),
        "tong_tin_nhan": db.query(Message).count(),
        "tong_tai_lieu": db.query(KnowledgeDocument).count(),
        "tong_bo_nho": db.query(MemoryItem).count(),
        "tong_dich_vu": db.query(ServiceItem).count(),
        "tong_luot_goi_ai": db.query(AIUsage).count(),
        "loi_goi_ai": db.query(AIUsage).filter(AIUsage.status == "loi").count(),
        "chi_phi_ai_uoc_tinh": sum(x.estimated_cost or 0 for x in db.query(AIUsage).all()),
        "nha_cung_cap_mac_dinh": (
            {"id": provider_mac_dinh.id, "name": provider_mac_dinh.name, "model": provider_mac_dinh.model,
             "status": provider_mac_dinh.status} if provider_mac_dinh else None
        ),
        "tai_khoan_zalo": tai_khoan_zalo,
    }


def _loc_theo_thoi_gian(query, cot, tu_ngay: str = "", den_ngay: str = ""):
    if tu_ngay:
        try:
            query = query.filter(cot >= datetime.fromisoformat(tu_ngay))
        except ValueError:
            pass
    if den_ngay:
        try:
            query = query.filter(cot <= datetime.fromisoformat(den_ngay))
        except ValueError:
            pass
    return query


def tinh_phan_tich(db: Session, tu_ngay: str = "", den_ngay: str = "") -> dict:
    q_msg = _loc_theo_thoi_gian(db.query(Message), Message.created_at, tu_ngay, den_ngay)
    total = q_msg.count()
    bot = q_msg.filter(Message.sender_type == "bot").count()
    khach = q_msg.filter(Message.sender_type == "khach").count()
    nhan_vien = q_msg.filter(Message.sender_type == "nhan_vien").count()
    q_ai = _loc_theo_thoi_gian(db.query(AIUsage), AIUsage.created_at, tu_ngay, den_ngay)
    ds_ai = q_ai.all()
    return {
        "tong_tin_nhan": total,
        "tin_nhan_khach": khach,
        "tin_nhan_bot": bot,
        "tin_nhan_nhan_vien": nhan_vien,
        "ty_le_bot": (bot / total if total else 0),
        "tong_hoi_thoai": db.query(Conversation).count(),
        "hoi_thoai_chuyen_nhan_vien": db.query(Conversation).filter(Conversation.human_takeover == True).count(),
        "luot_goi_ai": q_ai.count(),
        "loi_ai": q_ai.filter(AIUsage.status == "loi").count(),
        "tong_token_prompt": sum(x.prompt_tokens or 0 for x in ds_ai),
        "tong_token_completion": sum(x.completion_tokens or 0 for x in ds_ai),
        "chi_phi_uoc_tinh": sum(x.estimated_cost or 0 for x in ds_ai),
        "tong_khach_hang": db.query(Customer).count(),
        "tong_tai_lieu": db.query(KnowledgeDocument).count(),
        "tong_bo_nho": db.query(MemoryItem).count(),
        "tong_quy_tac": db.query(AutomationRule).count(),
        "tong_lan_chay_quy_tac": db.query(RuleRun).count(),
    }


@router.get("/api/phan-tich")
def analytics(tu_ngay: str = "", den_ngay: str = "", db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    return tinh_phan_tich(db, tu_ngay, den_ngay)


@router.get("/api/phan-tich/xuat-csv")
def xuat_csv(tu_ngay: str = "", den_ngay: str = "", db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    """Xuất báo cáo phân tích dạng CSV (tiếng Việt, UTF-8 có BOM cho Excel)."""
    so_lieu = tinh_phan_tich(db, tu_ngay, den_ngay)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Chi so", "Gia tri"])
    nhan = {
        "tong_tin_nhan": "Tổng tin nhắn", "tin_nhan_khach": "Tin nhắn khách gửi",
        "tin_nhan_bot": "Tin nhắn bot gửi", "tin_nhan_nhan_vien": "Tin nhắn nhân viên gửi",
        "ty_le_bot": "Tỷ lệ bot trả lời", "tong_hoi_thoai": "Tổng hội thoại",
        "hoi_thoai_chuyen_nhan_vien": "Hội thoại chuyển nhân viên",
        "luot_goi_ai": "Số lượt gọi AI", "loi_ai": "Số lượt gọi AI lỗi",
        "tong_token_prompt": "Tổng token đầu vào", "tong_token_completion": "Tổng token đầu ra",
        "chi_phi_uoc_tinh": "Chi phí AI ước tính", "tong_khach_hang": "Tổng khách hàng",
        "tong_tai_lieu": "Tổng tài liệu tri thức", "tong_bo_nho": "Tổng ký ức",
        "tong_quy_tac": "Tổng quy tắc tự động hóa", "tong_lan_chay_quy_tac": "Tổng lần chạy quy tắc",
    }
    for khoa, gia_tri in so_lieu.items():
        writer.writerow([nhan.get(khoa, khoa), gia_tri])
    writer.writerow([])
    writer.writerow(["Luot goi AI gan nhat (toi da 100)"])
    writer.writerow(["Thoi gian", "Mo hinh", "Token dau vao", "Token dau ra", "Chi phi uoc tinh", "Trang thai", "Loi"])
    for x in db.query(AIUsage).order_by(AIUsage.id.desc()).limit(100).all():
        writer.writerow([x.created_at, x.model, x.prompt_tokens, x.completion_tokens, x.estimated_cost, x.status, x.error or ""])
    noi_dung = "\ufeff" + buf.getvalue()
    ten_file = f"bao-cao-phan-tich-{datetime.now().strftime('%Y%m%d%H%M%S')}.csv"
    return StreamingResponse(
        io.BytesIO(noi_dung.encode("utf-8")),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{ten_file}"'},
    )


@router.get("/api/phan-tich/xuat-excel")
def xuat_excel(tu_ngay: str = "", den_ngay: str = "", db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    """Xuất báo cáo phân tích tổng hợp ra .xlsx (sheet Số liệu + sheet Lượt gọi AI)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    so_lieu = tinh_phan_tich(db, tu_ngay, den_ngay)
    nhan = {
        "tong_tin_nhan": "Tổng tin nhắn", "tin_nhan_khach": "Tin nhắn khách gửi",
        "tin_nhan_bot": "Tin nhắn bot gửi", "tin_nhan_nhan_vien": "Tin nhắn nhân viên gửi",
        "ty_le_bot": "Tỷ lệ bot trả lời", "tong_hoi_thoai": "Tổng hội thoại",
        "hoi_thoai_chuyen_nhan_vien": "Hội thoại chuyển nhân viên",
        "luot_goi_ai": "Số lượt gọi AI", "loi_ai": "Số lượt gọi AI lỗi",
        "tong_token_prompt": "Tổng token đầu vào", "tong_token_completion": "Tổng token đầu ra",
        "chi_phi_uoc_tinh": "Chi phí AI ước tính", "tong_khach_hang": "Tổng khách hàng",
        "tong_tai_lieu": "Tổng tài liệu tri thức", "tong_bo_nho": "Tổng ký ức",
        "tong_quy_tac": "Tổng quy tắc tự động hóa", "tong_lan_chay_quy_tac": "Tổng lần chạy quy tắc",
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Số liệu"
    ws.append(["Chỉ số", "Giá trị"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for khoa, gia_tri in so_lieu.items():
        ws.append([nhan.get(khoa, khoa), gia_tri])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20
    ws2 = wb.create_sheet("Lượt gọi AI")
    ws2.append(["Thời gian", "Mô hình", "Token đầu vào", "Token đầu ra", "Chi phí ước tính", "Trạng thái", "Lỗi"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for x in db.query(AIUsage).order_by(AIUsage.id.desc()).limit(500).all():
        ws2.append([x.created_at.strftime("%d/%m/%Y %H:%M") if x.created_at else "", x.model,
                    x.prompt_tokens, x.completion_tokens, x.estimated_cost, x.status, x.error or ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    ten_file = f"bao-cao-phan-tich-{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{ten_file}"'},
    )


@router.get("/api/nhat-ky")
def logs(
    muc_do: str = Query(default="", description="Lọc theo mức độ: info, warning, error"),
    module: str = Query(default="", description="Lọc theo phân hệ"),
    tu_khoa: str = Query(default="", description="Tìm trong nội dung nhật ký"),
    db: Session = Depends(get_db), u: User = Depends(get_current_user),
):
    q = db.query(SystemLog)
    if muc_do:
        q = q.filter(SystemLog.level == muc_do)
    if module:
        q = q.filter(SystemLog.module == module)
    if tu_khoa:
        q = q.filter(SystemLog.message.contains(tu_khoa))
    return [
        {"id": l.id, "level": l.level, "module": l.module, "message": l.message,
         "trace_id": l.trace_id, "created_at": str(l.created_at)}
        for l in q.order_by(SystemLog.id.desc()).limit(200).all()
    ]
